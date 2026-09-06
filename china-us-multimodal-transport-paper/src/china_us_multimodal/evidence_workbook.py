from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class EvidenceInput:
    category: str
    parameter: str
    value: float | str
    unit: str
    basis: str
    model_use: str
    source: str | None
    notes: str | None


REQUIRED_BASELINE_PARAMETERS = {
    "Shanghai → Los Angeles",
    "Shanghai → New York",
    "Average USEC premium ΔF",
    "LA ↔ Chicago intermodal",
    "Dry van all-in spot",
    "Symmetric handling baseline",
    "Normal α_Panama",
    "Restricted α_Panama",
    "Stress α_Panama",
    "LA / Inland Empire → Chicago",
}


def read_baseline_inputs(path: str | Path) -> dict[str, EvidenceInput]:
    workbook = load_workbook(Path(path), data_only=True, read_only=True)
    if "Baseline_Inputs" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the Baseline_Inputs sheet.")
    sheet = workbook["Baseline_Inputs"]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows) if row and row[0] == "Category"),
        None,
    )
    if header_index is None:
        raise ValueError("Baseline_Inputs has no recognizable header row.")

    result: dict[str, EvidenceInput] = {}
    for row in rows[header_index + 1 :]:
        if not row or row[0] is None:
            continue
        values = list(row) + [None] * (8 - len(row))
        item = EvidenceInput(
            category=str(values[0]),
            parameter=str(values[1]),
            value=values[2],
            unit=str(values[3]),
            basis=str(values[4]),
            model_use=str(values[5]),
            source=None if values[6] is None else str(values[6]),
            notes=None if values[7] is None else str(values[7]),
        )
        if item.parameter in result:
            raise ValueError(f"Duplicate baseline parameter: {item.parameter}")
        result[item.parameter] = item

    missing = sorted(REQUIRED_BASELINE_PARAMETERS - result.keys())
    if missing:
        raise ValueError(f"Missing required baseline parameters: {missing}")
    return result
