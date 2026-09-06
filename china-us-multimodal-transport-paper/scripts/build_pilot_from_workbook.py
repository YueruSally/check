from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict, deque
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


NODE_IDS = {
    "Xi'an": "xian",
    "Shanghai": "shanghai",
    "Ningbo": "ningbo",
    "Los Angeles/Long Beach": "la_lb",
    "Seattle/Tacoma": "seattle_tacoma",
    "New York/New Jersey": "ny_nj",
    "Chicago": "chicago",
    "Memphis": "memphis",
    "Columbus": "columbus",
}

NODE_KINDS = {
    "Xi'an": "cn_origin",
    "Shanghai": "cn_port",
    "Ningbo": "cn_port",
    "Los Angeles/Long Beach": "us_port",
    "Seattle/Tacoma": "us_port",
    "New York/New Jersey": "us_port",
    "Chicago": "destination",
    "Memphis": "destination",
    "Columbus": "destination",
}

CASE_SHEETS = {
    "D1": "Batches_Chicago",
    "D2": "Batches_Memphis",
    "D3": "Batches_Columbus",
}


class WorkbookValidationError(ValueError):
    pass


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _table(workbook, sheet_name: str, first_header: str) -> tuple[list[dict[str, Any]], int]:
    if sheet_name not in workbook.sheetnames:
        raise WorkbookValidationError(f"Missing required sheet: {sheet_name}")
    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    header_index = next(
        (index for index, row in enumerate(rows) if row and row[0] == first_header),
        None,
    )
    if header_index is None:
        raise WorkbookValidationError(
            f"{sheet_name} has no header beginning with {first_header!r}."
        )
    header = [str(value) if value is not None else "" for value in rows[header_index]]
    records: list[dict[str, Any]] = []
    for source_row, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not row or all(value in (None, "") for value in row):
            continue
        record = {key: value for key, value in zip(header, row) if key}
        record["_source_row"] = source_row
        records.append(record)
    return records, header_index + 1


def _write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _require_unique(records: list[dict[str, Any]], key: str, table: str, errors: list[str]) -> None:
    counts = Counter(record.get(key) for record in records)
    duplicates = sorted(str(value) for value, count in counts.items() if count > 1)
    if duplicates:
        errors.append(f"{table}: duplicate {key} values: {duplicates}")


def _reachable(adjacency: dict[str, set[str]], origin: str, destination: str) -> bool:
    queue = deque([origin])
    seen = {origin}
    while queue:
        node = queue.popleft()
        if node == destination:
            return True
        for next_node in adjacency.get(node, set()):
            if next_node not in seen:
                seen.add(next_node)
                queue.append(next_node)
    return False


def build(source: Path, output: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    errors: list[str] = []

    nodes_raw, nodes_header_row = _table(workbook, "Nodes", "EnglishName")
    arcs_raw, arcs_header_row = _table(workbook, "Arcs_All", "ArcID")
    timetable_raw, timetable_header_row = _table(workbook, "Timetable", "ServiceID")
    departures_raw, departures_header_row = _table(
        workbook, "Service_Departures_8wk", "ServiceID"
    )
    transfers_raw, transfers_header_row = _table(
        workbook, "Transshipment_Model", "Node"
    )
    batches_by_case: dict[str, list[dict[str, Any]]] = {}
    batch_header_rows: dict[str, int] = {}
    for case_id, sheet_name in CASE_SHEETS.items():
        batches_by_case[case_id], batch_header_rows[case_id] = _table(
            workbook, sheet_name, "BatchID"
        )

    _require_unique(nodes_raw, "EnglishName", "Nodes", errors)
    _require_unique(arcs_raw, "ArcID", "Arcs_All", errors)
    _require_unique(timetable_raw, "ServiceID", "Timetable", errors)
    if set(record["EnglishName"] for record in nodes_raw) != set(NODE_IDS):
        errors.append("Nodes do not exactly match the approved nine-node pilot set.")

    node_names = {record["EnglishName"] for record in nodes_raw}
    arc_by_signature: dict[tuple[str, str, str], dict[str, Any]] = {}
    adjacency: dict[str, set[str]] = defaultdict(set)
    for arc in arcs_raw:
        origin, destination = arc.get("OriginEN"), arc.get("DestEN")
        mode = str(arc.get("Mode", "")).strip().lower()
        if mode == "water":
            mode = "ocean"
        if origin not in node_names or destination not in node_names:
            errors.append(f"Arc {arc.get('ArcID')} has an unknown endpoint.")
        signature = (str(origin), str(destination), mode)
        if signature in arc_by_signature:
            errors.append(f"Duplicate arc signature: {signature}")
        arc_by_signature[signature] = arc
        adjacency[str(origin)].add(str(destination))
        for field in ("TransitTime_h", "TotalCost_USD_per_TEU", "Capacity_TEUday"):
            if arc.get(field) is None or float(arc[field]) <= 0:
                errors.append(f"Arc {arc.get('ArcID')} has invalid {field}.")

    services_by_arc: dict[str, list[str]] = defaultdict(list)
    service_meta: dict[str, dict[str, Any]] = {}
    for service in timetable_raw:
        mode = str(service.get("Mode", "")).strip().lower()
        if mode == "water":
            mode = "ocean"
        signature = (str(service.get("OriginEN")), str(service.get("DestEN")), mode)
        arc = arc_by_signature.get(signature)
        if arc is None:
            errors.append(f"Service {service.get('ServiceID')} has no matching arc {signature}.")
            continue
        service_id = str(service["ServiceID"])
        services_by_arc[str(arc["ArcID"])].append(service_id)
        service_meta[service_id] = service
        if float(service.get("Headway_Hours") or 0) <= 0:
            errors.append(f"Service {service_id} has a non-positive headway.")

    departures_by_service: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for departure in departures_raw:
        service_id = str(departure.get("ServiceID"))
        departures_by_service[service_id].append(departure)
        if service_id not in service_meta:
            errors.append(f"Departure references unknown service {service_id}.")
            continue
        expected_arrival = float(departure["DepartureHour"]) + float(service_meta[service_id]["time"])
        if abs(float(departure["ArrivalHour"]) - expected_arrival) > 1e-6:
            errors.append(
                f"{service_id} departure {departure.get('DepartureNo')} has inconsistent arrival."
            )

    for service_id, meta in service_meta.items():
        departures = sorted(
            departures_by_service.get(service_id, []),
            key=lambda item: float(item["DepartureHour"]),
        )
        expected_count = int(meta["Frequency_per_week"]) * 8
        if len(departures) != expected_count:
            errors.append(
                f"{service_id} has {len(departures)} departures; expected {expected_count}."
            )
        hours = [float(item["DepartureHour"]) for item in departures]
        if any(right <= left for left, right in zip(hours, hours[1:])):
            errors.append(f"{service_id} departure hours are not strictly increasing.")

    batch_totals: dict[str, float] = {}
    for case_id, records in batches_by_case.items():
        _require_unique(records, "BatchID", CASE_SHEETS[case_id], errors)
        destinations = {str(record.get("DestEN")) for record in records}
        if len(records) != 40:
            errors.append(f"{CASE_SHEETS[case_id]} must contain exactly 40 batches.")
        if len(destinations) != 1:
            errors.append(f"{CASE_SHEETS[case_id]} mixes destinations: {destinations}")
        for record in records:
            if str(record.get("OriginEN")) != "Xi'an":
                errors.append(f"{CASE_SHEETS[case_id]} has a non-Xi'an origin.")
            if float(record.get("QuantityTEU") or 0) <= 0:
                errors.append(f"{CASE_SHEETS[case_id]} has a non-positive quantity.")
            if float(record.get("LT") or -1) < float(record.get("ET") or 0):
                errors.append(f"{CASE_SHEETS[case_id]} has LT before ET.")
            if not _reachable(adjacency, str(record["OriginEN"]), str(record["DestEN"])):
                errors.append(
                    f"Batch {case_id}-{record['BatchID']} has no topological route."
                )
        batch_totals[case_id] = sum(float(record["QuantityTEU"]) for record in records)

    valid_transfers: list[dict[str, Any]] = []
    transfer_keys: set[tuple[str, str, str]] = set()
    for transfer in transfers_raw:
        key = (
            str(transfer.get("Node")),
            str(transfer.get("FromMode", "")).lower().replace("water", "ocean"),
            str(transfer.get("ToMode", "")).lower().replace("water", "ocean"),
        )
        if key in transfer_keys:
            errors.append(f"Duplicate transfer rule: {key}")
        transfer_keys.add(key)
        if key[0] not in node_names:
            errors.append(f"Transfer rule uses unknown node: {key[0]}")
        if float(transfer.get("TransferCost_USD_per_TEU") or -1) < 0:
            errors.append(f"Transfer rule {key} has a negative cost.")
        if float(transfer.get("TransferTime_h") or -1) < 0:
            errors.append(f"Transfer rule {key} has a negative duration.")
        valid_transfers.append(transfer)

    if errors:
        raise WorkbookValidationError("\n".join(errors))

    output.mkdir(parents=True, exist_ok=True)
    nodes = []
    for record in nodes_raw:
        name = str(record["EnglishName"])
        nodes.append(
            {
                "id": NODE_IDS[name],
                "name": name,
                "kind": NODE_KINDS[name],
                "country": "CN" if record["Country"] == "China" else "US",
                "capacity_feu_per_day": float(record["NodeCap_TEUday"]) / 2,
                "source_sheet": "Nodes",
                "source_row": record["_source_row"],
            }
        )

    arcs = []
    arc_id_map: dict[str, str] = {}
    for record in arcs_raw:
        source_arc_id = str(record["ArcID"])
        arc_id = source_arc_id.lower()
        arc_id_map[source_arc_id] = arc_id
        mode = str(record["Mode"]).lower().replace("water", "ocean")
        services = services_by_arc.get(source_arc_id, [])
        headways = [float(service_meta[item]["Headway_Hours"]) for item in services]
        arcs.append(
            {
                "id": arc_id,
                "origin": NODE_IDS[str(record["OriginEN"])],
                "destination": NODE_IDS[str(record["DestEN"])],
                "mode": mode,
                "transit_time_h": float(record["TransitTime_h"]),
                "cost_usd_per_feu": float(record["TotalCost_USD_per_TEU"]) * 2,
                "distance_km": float(record["Distance_km"]),
                "capacity_feu_per_day": float(record["Capacity_TEUday"]) / 2,
                "service_ids": "|".join(services),
                "planning_wait_h": min(headways) / 2 if headways else 0,
                "via_panama": int(bool(record["PanamaExposed"])),
                "source_sheet": "Arcs_All",
                "source_row": record["_source_row"],
            }
        )

    services = []
    for service_id, record in service_meta.items():
        mode = str(record["Mode"]).lower().replace("water", "ocean")
        source_arc = arc_by_signature[
            (str(record["OriginEN"]), str(record["DestEN"]), mode)
        ]
        departures = sorted(
            float(item["DepartureHour"]) for item in departures_by_service[service_id]
        )
        daily_capacity_feu = float(record["Capacity_TEUday"]) / 2
        headway_h = float(record["Headway_Hours"])
        services.append(
            {
                "id": service_id.lower(),
                "arc_id": arc_id_map[str(source_arc["ArcID"])],
                "departures_h": "|".join(f"{value:g}" for value in departures),
                "repeat_every_h": 8 * 7 * 24,
                "capacity_feu": daily_capacity_feu * headway_h / 24,
                "cutoff_h": 0,
                "source_sheet": "Service_Departures_8wk",
                "source_rows": f"3:{2 + len(departures_raw)}",
            }
        )
    service_id_map = {item["id"].upper(): item["id"] for item in services}
    for arc in arcs:
        arc["service_ids"] = "|".join(
            service_id_map[item] for item in arc["service_ids"].split("|") if item
        )

    transfers = []
    for record in valid_transfers:
        transfers.append(
            {
                "node_id": NODE_IDS[str(record["Node"])],
                "from_mode": str(record["FromMode"]).lower().replace("water", "ocean"),
                "to_mode": str(record["ToMode"]).lower().replace("water", "ocean"),
                "duration_h": float(record["TransferTime_h"]),
                "cost_usd_per_feu": float(record["TransferCost_USD_per_TEU"]) * 2,
                "source_sheet": "Transshipment_Model",
                "source_row": record["_source_row"],
            }
        )

    shipments = []
    for case_id, records in batches_by_case.items():
        sheet_name = CASE_SHEETS[case_id]
        for record in records:
            shipments.append(
                {
                    "id": f"{case_id}-{int(record['BatchID']):02d}",
                    "origin": NODE_IDS[str(record["OriginEN"])],
                    "destination": NODE_IDS[str(record["DestEN"])],
                    "quantity_feu": float(record["QuantityTEU"]) / 2,
                    "release_h": float(record["ET"]),
                    "due_h": float(record["LT"]),
                    "tardiness_usd_per_feu_h": float(record["PenaltyCost_per_TEU_h"]) * 2,
                    "case_id": case_id,
                    "source_batch_id": int(record["BatchID"]),
                    "source_sheet": sheet_name,
                    "source_row": record["_source_row"],
                }
            )

    _write_csv(
        output / "nodes.csv",
        ["id", "name", "kind", "country", "capacity_feu_per_day", "source_sheet", "source_row"],
        nodes,
    )
    _write_csv(
        output / "arcs.csv",
        [
            "id", "origin", "destination", "mode", "transit_time_h",
            "cost_usd_per_feu", "distance_km", "capacity_feu_per_day",
            "service_ids", "planning_wait_h", "via_panama", "source_sheet", "source_row",
        ],
        arcs,
    )
    _write_csv(
        output / "services.csv",
        [
            "id", "arc_id", "departures_h", "repeat_every_h", "capacity_feu",
            "cutoff_h", "source_sheet", "source_rows",
        ],
        services,
    )
    _write_csv(
        output / "shipments.csv",
        [
            "id", "origin", "destination", "quantity_feu", "release_h", "due_h",
            "tardiness_usd_per_feu_h", "case_id", "source_batch_id", "source_sheet", "source_row",
        ],
        shipments,
    )
    _write_csv(
        output / "transfers.csv",
        [
            "node_id", "from_mode", "to_mode", "duration_h", "cost_usd_per_feu",
            "source_sheet", "source_row",
        ],
        transfers,
    )

    report = {
        "source_workbook": source.name,
        "source_sha256": _sha256(source),
        "output_quantity_unit": "FEU",
        "conversion_rules": {
            "quantity_and_capacity": "TEU / 2 = FEU",
            "per_unit_cost_and_penalty": "USD/TEU * 2 = USD/FEU",
            "service_capacity": "daily FEU capacity * headway hours / 24",
            "calendar": "explicit 8-week departure hours; repeats every 1344 hours",
        },
        "counts": {
            "nodes": len(nodes),
            "arcs": len(arcs),
            "services": len(services),
            "departures": len(departures_raw),
            "transfer_rules": len(transfers),
            "shipments": len(shipments),
        },
        "cases": {
            case_id: {
                "sheet": CASE_SHEETS[case_id],
                "shipments": len(batches_by_case[case_id]),
                "total_teu": batch_totals[case_id],
                "total_feu": batch_totals[case_id] / 2,
            }
            for case_id in CASE_SHEETS
        },
        "checks": {
            "unique_ids": True,
            "known_arc_endpoints": True,
            "service_arc_matches": True,
            "departure_counts_and_arrivals": True,
            "batch_reachability": True,
            "nonnegative_costs_times_capacities": True,
        },
        "source_header_rows": {
            "Nodes": nodes_header_row,
            "Arcs_All": arcs_header_row,
            "Timetable": timetable_header_row,
            "Service_Departures_8wk": departures_header_row,
            "Transshipment_Model": transfers_header_row,
            **{CASE_SHEETS[key]: value for key, value in batch_header_rows.items()},
        },
    }
    (output / "conversion_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Validate and convert the model-ready China-US workbook to FEU CSV tables."
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    report = build(args.workbook, args.output)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
